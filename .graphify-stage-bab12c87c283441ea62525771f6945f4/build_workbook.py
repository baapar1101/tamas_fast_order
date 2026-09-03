import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

def build():
    wb = Workbook()
    
    # Define styles
    navy = "00768F"
    border = Side(style="thin", color="D9E1EA")
    hf = Font(name="Vazirmatn", size=11, bold=True, color="FFFFFF")
    bf = Font(name="Vazirmatn", size=11)
    
    def setup(ws, heads, widths):
        ws.delete_rows(1, ws.max_row)
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.showGridLines = True
        ws.append(heads)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(1, len(heads)).coordinate}"
        for i in range(1, len(heads) + 1):
            c = ws.cell(1, i)
            c.font = hf
            c.fill = PatternFill("solid", fgColor=navy)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(bottom=border)
            ws.column_dimensions[c.column_letter].width = widths[i - 1]
        ws.row_dimensions[1].height = 30

    # 1. Categories sheet
    ws_cat = wb.active
    ws_cat.title = "Categories"
    setup(ws_cat, ["category_name", "category_fa_name", "filter_name", "icon"], [24, 26, 24, 18])
    categories_data = [
        ["Mobile phones", "گوشی موبایل", "گوشی", "mobile.svg"],
        ["Tablets", "تبلت", "تبلت", "mobile.svg"],
        ["Earphones & AirPods", "هدفون و هندزفری", "هندزفری", "airpod.svg"],
        ["Powerbanks", "پاوربانک", "پاوربانک", "powerbank.svg"],
        ["Chargers & Adapters", "شارژر و آداپتور", "شارژر", "adapter.svg"],
        ["Cables", "کابل و تبدیل", "کابل", "cable.svg"],
        ["Covers & Cases", "قاب و کاور", "کاور", "cover.svg"],
        ["Screen Protectors", "محافظ صفحه", "گلس", "glass.svg"],
        ["Speakers", "اسپیکر", "اسپیکر", "speaker.svg"],
        ["Smartwatches", "ساعت و مچ‌بند", "ساعت", "watch.svg"],
        ["Car Accessories", "لوازم جانبی خودرو", "خودرو", "caracsories.svg"],
        ["Holders & Mounts", "پایه نگهدارنده", "هولدر", "holder.svg"],
        ["Mouse & Keyboard", "ماوس و کیبورد", "ماوس", "mouse_keyboard.svg"],
        ["Gaming Accessories", "لوازم گیمینگ", "گیمینگ", "gaming.svg"],
        ["Beauty & Health", "محصولات زیبایی و سلامت", "سلامت", "health.svg"],
        ["Tools & Hardware", "ابزار و تجهیزات", "ابزار", "tools.svg"]
    ]
    for r in categories_data:
        ws_cat.append(r)

    # 2. Brands sheet
    ws_brand = wb.create_sheet("Brands")
    setup(ws_brand, ["brand_name", "brand_fa_name", "icon"], [22, 24, 18])
    brands_data = [
        ["Apple", "اپل", "apple.svg"],
        ["Samsung", "سامسونگ", "samsung.svg"],
        ["Xiaomi", "شیائومی", "xiaomi.svg"],
        ["Anker", "انکر", "anker.svg"],
        ["Hoco", "هوکو", "hoco.svg"],
        ["Macdodo", "مک‌دودو", "macdodo.svg"],
        ["Yesido", "یسیدو", "yesido.svg"],
        ["QCY", "کیو‌سی‌وای", "qcy.svg"],
        ["Oxygen", "اکسیژن", "oxygen.svg"],
        ["Hadron", "هادرون", "hadron.svg"],
        ["Borofone", "برافون", "borofone.svg"],
        ["Haylou", "هایلو", "haylou.svg"],
        ["Nokia", "نوکیا", "nokia.svg"],
        ["Redmi", "ردمی", "redmi.svg"],
        ["Honor", "آنر", "honor.svg"],
        ["Baiko", "بایکو", "baiko.svg"],
        ["Earldom", "ارلدام", "earldom.svg"]
    ]
    for r in brands_data:
        ws_brand.append(r)

    # 3. Colors sheet
    ws_color = wb.create_sheet("Colors")
    setup(ws_color, ["color_name", "color_fa_name"], [22, 24])
    colors_data = [
        ["Black", "مشکی"],
        ["White", "سفید"],
        ["Blue", "آبی"],
        ["Deep Blue", "سرمه‌ای"],
        ["Cosmic Orange", "نارنجی"],
        ["Green", "سبز"],
        ["Gold", "طلایی"],
        ["Silver", "نقره‌ای"],
        ["Titanium Gray", "خاکستری تیتانیوم"],
        ["Purple", "بنفش"]
    ]
    for r in colors_data:
        ws_color.append(r)

    # 4. Products sheet
    ws_prod = wb.create_sheet("Products")
    heads = [
        "product_id", "Category", "Brand", "title", "model", "color", "color_en", "sku", 
        "price", "old_price", "discount", "stock", "warranty", "promotion", "status", 
        "image_url", "attribute_key", "attribute_value"
    ]
    setup(ws_prod, heads, [15, 18, 16, 36, 30, 22, 18, 18, 16, 16, 12, 10, 24, 12, 16, 45, 18, 24])

    sample_products = [
        [
            "P101", "گوشی موبایل", "اپل", "Apple iPhone 17 Pro Max 512GB", "iPhone 17 Pro Max", 
            "نارنجی (Cosmic Orange)", "Cosmic Orange", "APL-IP17PM-512-ORG", 450000000, 455000000, 
            1, 5, "18 ماه گارانتی شرکتی", True, "Active", 
            "https://images.unsplash.com/photo-1592899677977-9c10ca588bbd8?auto=format&fit=crop&w=400&q=80", 
            "حافظه داخلی", "512 گیگابایت"
        ],
        [
            "P102", "گوشی موبایل", "اپل", "Apple iPhone 17 Pro Max 512GB", "iPhone 17 Pro Max", 
            "سرمه‌ای (Deep Blue)", "Deep Blue", "APL-IP17PM-512-BLU", 477900000, 482000000, 
            1, 3, "18 ماه گارانتی شرکتی", True, "Active", 
            "https://images.unsplash.com/photo-1592286927505-2fd0f9d7c6f4?auto=format&fit=crop&w=400&q=80", 
            "حافظه داخلی", "512 گیگابایت"
        ],
        [
            "P103", "گوشی موبایل", "سامسونگ", "Samsung Galaxy S25 Ultra 512GB", "Galaxy S25 Ultra", 
            "مشکی (Black)", "Black", "SAM-S25U-512-BLK", 385000000, 395000000, 
            3, 8, "18 ماه گارانتی داریا همراه", True, "Active", 
            "https://images.unsplash.com/photo-1610945265064-0e34e5519bbf?auto=format&fit=crop&w=400&q=80", 
            "حافظه رم", "12 گیگابایت"
        ],
        [
            "P104", "هدفون و هندزفری", "هوکو", "Hoco EW04 TWS Wireless Earphones", "Hoco EW04", 
            "سفید (White)", "White", "HOC-EW04-WHT", 850000, 950000, 
            10, 25, "6 ماه گارانتی پژواک", False, "Active", 
            "https://images.unsplash.com/photo-1590658268037-6bf12165a8df?auto=format&fit=crop&w=400&q=80", 
            "نوع اتصال", "بلوتوث 5.3"
        ],
        [
            "P105", "پاوربانک", "انکر", "Anker 335 Power Bank 20000mAh 20W", "Anker 335 20K", 
            "مشکی (Black)", "Black", "ANK-A1363-BLK", 2450000, 2700000, 
            9, 15, "18 ماه گارانتی ایستا", True, "Active", 
            "https://images.unsplash.com/photo-1609592424074-9f44cf2f01f0?auto=format&fit=crop&w=400&q=80", 
            "ظرفیت باتری", "20000 میلی‌آمپر"
        ]
    ]

    for p in sample_products:
        ws_prod.append(p)

    # Format cells in Products
    for row in ws_prod.iter_rows(min_row=2, max_row=ws_prod.max_row):
        for c in row:
            c.font = bf
            c.alignment = Alignment(vertical="center", horizontal="center" if c.column in [1, 9, 10, 11, 12, 14, 15] else "right", wrap_text=True)
            c.border = Border(bottom=border)
        row[8].number_format = '#,##0'
        row[9].number_format = '#,##0'

    # Data Validations
    for formula, rng in [
        ("=Categories!$B$2:$B$100", "B2:B1000"),
        ("=Brands!$B$2:$B$100", "C2:C1000"),
        ("=Colors!$B$2:$B$100", "F2:F1000"),
        ('"بدون گارانتی,6 ماهه,12 ماهه,18 ماهه,24 ماهه"', "M2:M1000"),
        ('"TRUE,FALSE"', "N2:N1000"),
        ('"Active,Inactive,Draft"', "O2:O1000")
    ]:
        dv = DataValidation(type="list", formula1=formula, allow_blank=True)
        ws_prod.add_data_validation(dv)
        dv.add(rng)

    # Out of stock highlight
    ws_prod.conditional_formatting.add("L2:L1000", CellIsRule(operator="lessThan", formula=["1"], fill=PatternFill("solid", fgColor="FDE2E2")))

    # Add Table formatting
    tab = Table(displayName="ProductsTable", ref=f"A1:R{max(ws_prod.max_row, 2)}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    ws_prod.add_table(tab)

    # 5. Orders sheet
    ws_orders = wb.create_sheet("Orders")
    setup(ws_orders, ["order_id", "date", "customer_name", "phone", "address", "items_json", "total_price", "status"], [22, 20, 24, 16, 35, 45, 18, 14])
    ws_orders.sheet_view.rightToLeft = True

    # 6. Settings sheet
    ws_settings = wb.create_sheet("Settings")
    setup(ws_settings, ["key", "value"], [25, 45])
    ws_settings.sheet_view.rightToLeft = True
    settings_data = [
        ["store_name", "تماس مارکت (KarPlus)"],
        ["store_phone", "09135006644"],
        ["store_address", "کرمان، خیابان شهید نامجو، بعد از کوچه ۹، پلاک ۱۰۹"],
        ["min_wholesale_total", "5000000"]
    ]
    for r in settings_data:
        ws_settings.append(r)

    wb.save("fast_order_catalog.xlsx")
    print("Successfully built fast_order_catalog.xlsx")

if __name__ == "__main__":
    build()
