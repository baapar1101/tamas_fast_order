#!/usr/bin/env python3
"""
Google Sheets Synchronization Script for KarPlus / Tamas Fast Order
Spreadsheet ID: 1VjRRhxuLrr1HiEjLBNRrxXrZmk_pPjfu-s6xZY6bha4
Apps Script ID: 19Agigyvzwp-ji58rPI8ir10tHkf3qcWaxQIDtzBaDoMwUjAbTGWD0SAP

Provides synchronization capabilities between:
1. Local Excel workbook (fast_order_catalog.xlsx)
2. Google Spreadsheet via Google Apps Script Web API / Google Sheets CSV Export
3. Local catalog_backup.json
"""

import argparse
import csv
import json
import os
import sys
import time
from datetime import datetime, date
from urllib.parse import quote
import requests
import openpyxl

# Force UTF-8 output on Windows terminal
if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
if hasattr(sys.stderr, "reconfigure"):
    try:
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

DEFAULT_SPREADSHEET_ID = "1VjRRhxuLrr1HiEjLBNRrxXrZmk_pPjfu-s6xZY6bha4"
DEFAULT_EXCEL_FILE = "fast_order_catalog.xlsx"
DEFAULT_BACKUP_FILE = "catalog_backup.json"
CONFIG_FILE = "sync_config.json"

SHEET_NAMES = ["Categories", "Brands", "Colors", "Products", "Orders", "Settings"]


def load_config():
    """Load configuration from sync_config.json if it exists."""
    config = {
        "spreadsheet_id": DEFAULT_SPREADSHEET_ID,
        "web_app_url": os.environ.get("GAS_WEB_APP_URL", ""),
        # Matches the `admin_key` row in the Settings sheet. Required only once
        # that row exists, which locks down the users/orders/exportAll actions.
        "web_app_key": os.environ.get("GAS_WEB_APP_KEY", ""),
        "excel_file": DEFAULT_EXCEL_FILE,
        "backup_file": DEFAULT_BACKUP_FILE
    }
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r", encoding="utf-8") as f:
                saved = json.load(f)
                config.update(saved)
        except Exception as e:
            print(f"⚠️ Warning loading {CONFIG_FILE}: {e}")
    return config


def save_config(config):
    """Save current configuration to sync_config.json."""
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)


def read_excel_catalog(excel_path=DEFAULT_EXCEL_FILE):
    """Read local Excel workbook and return dictionary of sheets with rows."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    catalog = {}

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            catalog[sheet_name.lower()] = []
            continue

        headers = [str(cell).strip() if cell is not None else f"col_{i}" for i, cell in enumerate(rows[0])]
        sheet_data = []

        for r in rows[1:]:
            if not any(c is not None and str(c).strip() != "" for c in r):
                continue
            row_dict = {}
            for i, h in enumerate(headers):
                if i < len(r):
                    val = r[i]
                    if isinstance(val, (datetime, date)):
                        val = val.isoformat()
                    row_dict[h] = val
                else:
                    row_dict[h] = ""
            sheet_data.append(row_dict)

        catalog[sheet_name] = sheet_data

    # Map settings if Settings sheet exists
    settings_dict = {}
    if "Settings" in catalog:
        for row in catalog["Settings"]:
            k = str(row.get("key", "")).strip()
            if k:
                settings_dict[k] = row.get("value", "")

    return {
        "categories": catalog.get("Categories", []),
        "brands": catalog.get("Brands", []),
        "colors": catalog.get("Colors", []),
        "products": catalog.get("Products", []),
        "orders": catalog.get("Orders", []),
        "settings": settings_dict if settings_dict else catalog.get("Settings", [])
    }


def write_excel_catalog(catalog_data, output_path=DEFAULT_EXCEL_FILE):
    """Write catalog data (dict of lists) to Excel workbook."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default initial sheet

    for sheet_title in ["Categories", "Brands", "Colors", "Products", "Orders", "Settings"]:
        key = sheet_title.lower()
        rows = catalog_data.get(key, catalog_data.get(sheet_title, []))

        ws = wb.create_sheet(title=sheet_title)
        ws.sheet_view.rightToLeft = True

        if sheet_title == "Settings" and isinstance(rows, dict):
            rows = [{"key": k, "value": v} for k, v in rows.items()]

        if not rows:
            ws.append(["id"])
            continue

        if isinstance(rows[0], dict):
            headers = list(rows[0].keys())
            ws.append(headers)
            for item in rows:
                ws.append([item.get(h, "") for h in headers])
        elif isinstance(rows[0], list):
            for r in rows:
                ws.append(r)

    wb.save(output_path)
    print(f"✅ Saved catalog to Excel file: {output_path}")


def validate_catalog_images(catalog_data):
    """Reject obviously corrupted snapshots where one image was copied to nearly every product."""
    products = catalog_data.get("products", [])
    image_urls = [
        str(product.get("image_url", "")).strip()
        for product in products
        if str(product.get("image_url", "")).strip()
    ]
    if len(image_urls) < 20:
        return

    unique_count = len(set(image_urls))
    most_common_count = max(image_urls.count(url) for url in set(image_urls))
    dominant_ratio = most_common_count / len(image_urls)
    if unique_count == 1 or dominant_ratio >= 0.9:
        raise ValueError(
            "Image validation failed: the catalogue snapshot assigns the same "
            f"image to {most_common_count}/{len(image_urls)} products."
        )


def write_json_backup(catalog_data, output_path=DEFAULT_BACKUP_FILE):
    validate_catalog_images(catalog_data)
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(catalog_data, f, ensure_ascii=False, indent=2)
    print(f"✅ Saved local JSON backup: {output_path}")


def fetch_csv_from_gss(spreadsheet_id, sheet_name):
    """Fetch sheet content via Google Sheets CSV export URL (no API key required if accessible)."""
    url = f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/gviz/tq?tqx=out:csv&sheet={sheet_name}"
    resp = requests.get(url, timeout=10)
    if resp.status_code != 200:
        raise Exception(f"Failed to fetch CSV for sheet '{sheet_name}': HTTP {resp.status_code}")

    reader = csv.reader(resp.text.splitlines())
    rows = list(reader)
    if not rows:
        return []

    headers = [h.strip() for h in rows[0]]
    result = []
    for r in rows[1:]:
        if not any(cell.strip() for cell in r):
            continue
        row_dict = {}
        for i, h in enumerate(headers):
            row_dict[h] = r[i] if i < len(r) else ""
        result.append(row_dict)
    return result


def fetch_from_web_app(web_app_url, action="exportAll", key=""):
    """Fetch data from deployed Google Apps Script Web App endpoint."""
    url = f"{web_app_url}?action={action}"
    if key:
        url += f"&key={quote(key)}"
    resp = requests.get(url, timeout=15)
    resp.raise_for_status()
    data = resp.json()
    if not data.get("ok"):
        raise Exception(f"Apps Script error: {data.get('error', 'Unknown error')}")
    return data


def push_to_web_app(web_app_url, catalog_data):
    """Upload catalog data to Google Apps Script Web App."""
    payload = {
        "action": "syncCatalog",
        "catalog": catalog_data
    }
    headers = {"Content-Type": "application/json"}
    resp = requests.post(web_app_url, data=json.dumps(payload), headers=headers, timeout=30)
    resp.raise_for_status()
    data = resp.json()
    if not data.get("ok"):
        raise Exception(f"Apps Script push error: {data.get('error', 'Unknown error')}")
    return data


def cmd_pull(config, args):
    """Pull data from Google Sheets / Apps Script and update local Excel and JSON backup."""
    print("⬇️ Pulling catalog from Google Sheets...")

    data = None
    if config["web_app_url"]:
        try:
            print(f"Connecting to Web App URL: {config['web_app_url']}")
            res = fetch_from_web_app(config["web_app_url"], action="exportAll", key=config.get("web_app_key", ""))
            data = res.get("catalog", res)
            print("✅ Successfully pulled catalog via Web API")
        except Exception as e:
            print(f"⚠️ Web App fetch failed: {e}. Falling back to CSV export...")

    if not data:
        print(f"Fetching public CSV sheets from Spreadsheet ID: {config['spreadsheet_id']}")
        data = {
            "categories": fetch_csv_from_gss(config["spreadsheet_id"], "Categories"),
            "brands": fetch_csv_from_gss(config["spreadsheet_id"], "Brands"),
            "colors": fetch_csv_from_gss(config["spreadsheet_id"], "Colors"),
            "products": fetch_csv_from_gss(config["spreadsheet_id"], "Products"),
            "orders": fetch_csv_from_gss(config["spreadsheet_id"], "Orders"),
            "settings": fetch_csv_from_gss(config["spreadsheet_id"], "Settings")
        }

    # Save to local backup JSON only after a basic data-integrity check.
    write_json_backup(data, config["backup_file"])

    # Write to local Excel workbook
    write_excel_catalog(data, output_path=config["excel_file"])


def cmd_push(config, args):
    """Push local Excel data to Google Sheets via Apps Script Web App."""
    if not config["web_app_url"]:
        print("❌ Error: Web App URL is required for pushing data.")
        print("Please set Web App URL using: python sync_spreadsheet.py --set-url <YOUR_WEB_APP_URL>")
        sys.exit(1)

    print(f"⬆️ Reading local Excel workbook '{config['excel_file']}'...")
    catalog = read_excel_catalog(config["excel_file"])

    print(f"Uploading catalog ({len(catalog.get('products', []))} products, {len(catalog.get('categories', []))} categories)...")
    res = push_to_web_app(config["web_app_url"], catalog)
    print(f"✅ Sync complete! Server response: {res.get('message', 'Success')}")


def cmd_backup(config, args):
    """Rebuild the local JSON fallback from the local Excel workbook only."""
    print(f"📦 Reading local Excel workbook '{config['excel_file']}'...")
    catalog = read_excel_catalog(config["excel_file"])
    write_json_backup(catalog, config["backup_file"])


def cmd_diff(config, args):
    """Compare local Excel data against remote Google Sheet data."""
    print("🔍 Comparing local Excel catalog with Google Sheets catalog...")
    local_cat = read_excel_catalog(config["excel_file"])

    remote_cat = None
    if config["web_app_url"]:
        try:
            res = fetch_from_web_app(config["web_app_url"], action="catalog")
            remote_cat = res
        except Exception as e:
            print(f"⚠️ Web App fetch failed: {e}. Falling back to CSV export...")

    if not remote_cat:
        try:
            remote_cat = {
                "categories": fetch_csv_from_gss(config["spreadsheet_id"], "Categories"),
                "brands": fetch_csv_from_gss(config["spreadsheet_id"], "Brands"),
                "colors": fetch_csv_from_gss(config["spreadsheet_id"], "Colors"),
                "products": fetch_csv_from_gss(config["spreadsheet_id"], "Products"),
                "settings": fetch_csv_from_gss(config["spreadsheet_id"], "Settings")
            }
        except Exception as e:
            print(f"⚠️ Direct Google Sheets CSV fetch warning: {e}")
            remote_cat = {"categories": [], "brands": [], "colors": [], "products": [], "settings": {}}

    print("\n--- Summary Comparison ---")
    for key in ["categories", "brands", "colors", "products"]:
        l_cnt = len(local_cat.get(key, []))
        r_cnt = len(remote_cat.get(key, []))
        diff = l_cnt - r_cnt
        status = "Match" if diff == 0 else f"{'+' if diff > 0 else ''}{diff} difference"
        print(f"• {key.capitalize()}: Local={l_cnt} | Remote={r_cnt} -> {status}")

    # Compare products stock and price mismatches
    local_prods = {p.get("product_id") or p.get("sku"): p for p in local_cat.get("products", [])}
    remote_prods = {p.get("product_id") or p.get("sku"): p for p in remote_cat.get("products", [])}

    mismatches = []
    for p_id, lp in local_prods.items():
        if not p_id:
            continue
        rp = remote_prods.get(p_id)
        if not rp:
            mismatches.append(f"Product {p_id} ({lp.get('title')}) exists locally but missing in remote.")
        else:
            if str(lp.get("price")) != str(rp.get("price")):
                mismatches.append(f"Price diff for {p_id}: Local={lp.get('price')} vs Remote={rp.get('price')}")
            if str(lp.get("stock")) != str(rp.get("stock")):
                mismatches.append(f"Stock diff for {p_id}: Local={lp.get('stock')} vs Remote={rp.get('stock')}")

    if mismatches:
        print("\n--- Detailed Discrepancies ---")
        for m in mismatches[:15]:
            print(f"  ⚠️ {m}")
        if len(mismatches) > 15:
            print(f"  ... and {len(mismatches) - 15} more differences.")
    else:
        print("\n✨ All product prices and stock levels match perfectly!")


def cmd_watch(config, args):
    """Watch local Excel file for modifications and push automatically."""
    print(f"👀 Watching '{config['excel_file']}' for changes... (Press Ctrl+C to stop)")
    last_mtime = 0

    while True:
        try:
            if os.path.exists(config["excel_file"]):
                mtime = os.path.getmtime(config["excel_file"])
                if last_mtime != 0 and mtime > last_mtime:
                    print(f"\n🔄 File change detected at {datetime.now().strftime('%H:%M:%S')}. Triggering sync...")
                    try:
                        cmd_push(config, args)
                    except Exception as e:
                        print(f"❌ Sync failed: {e}")
                last_mtime = mtime
            time.sleep(2)
        except KeyboardInterrupt:
            print("\nStopped watcher.")
            break


def main():
    config = load_config()

    parser = argparse.ArgumentParser(
        description="Synchronization script for KarPlus / Tamas Fast Order Google Spreadsheet."
    )
    parser.add_argument("--pull", action="store_true", help="Pull data from Google Sheets to local Excel & JSON")
    parser.add_argument("--push", action="store_true", help="Push local Excel catalog to Google Sheets")
    parser.add_argument("--diff", action="store_true", help="Check differences between local and Google Sheets")
    parser.add_argument("--watch", action="store_true", help="Auto-sync when local Excel file changes")
    parser.add_argument("--backup", action="store_true", help="Build JSON backup from the local Excel file")
    parser.add_argument("--set-url", type=str, help="Set deployed Google Apps Script Web App URL")
    parser.add_argument("--sheet-id", type=str, help="Set Google Spreadsheet ID")
    parser.add_argument("--file", type=str, help="Specify local Excel file path")

    args = parser.parse_args()

    if args.set_url:
        config["web_app_url"] = args.set_url
        save_config(config)
        print(f"✅ Set Web App URL to: {config['web_app_url']}")

    if args.sheet_id:
        config["spreadsheet_id"] = args.sheet_id
        save_config(config)
        print(f"✅ Set Spreadsheet ID to: {config['spreadsheet_id']}")

    if args.file:
        config["excel_file"] = args.file

    if args.pull:
        cmd_pull(config, args)
    elif args.backup:
        cmd_backup(config, args)
    elif args.push:
        cmd_push(config, args)
    elif args.diff:
        cmd_diff(config, args)
    elif args.watch:
        cmd_watch(config, args)
    else:
        print("=== KarPlus Google Sheets Synchronization Tool ===")
        print(f"• Spreadsheet ID: {config['spreadsheet_id']}")
        print(f"• Local Excel File: {config['excel_file']}")
        print(f"• Web App URL: {config['web_app_url'] or '(Not configured yet)'}")
        print("\nUsage examples:")
        print("  python sync_spreadsheet.py --diff               # Compare local vs Google Sheets")
        print("  python sync_spreadsheet.py --pull               # Download Google Sheets to Excel")
        print("  python sync_spreadsheet.py --set-url <WEB_URL>  # Configure Web App URL")
        print("  python sync_spreadsheet.py --push               # Upload local Excel to Google Sheets")
        print("  python sync_spreadsheet.py --watch              # Auto-push on Excel file edit")
        print("-------------------------------------------------")
        cmd_diff(config, args)


if __name__ == "__main__":
    main()
