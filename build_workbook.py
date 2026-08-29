from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule

wb=load_workbook("karplus_google_sheet_template.xlsx")
navy="17324D"; border=Side(style="thin",color="D9E1EA")
hf=Font(name="Vazirmatn",size=11,bold=True,color="FFFFFF"); bf=Font(name="Vazirmatn",size=11)
def setup(ws,heads,widths):
    ws.delete_rows(1,ws.max_row); ws.sheet_view.rightToLeft=True; ws.sheet_view.showGridLines=False
    ws.append(heads); ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{ws.cell(1,len(heads)).coordinate}"
    for i in range(1,len(heads)+1):
        c=ws.cell(1,i); c.font=hf; c.fill=PatternFill("solid",fgColor=navy); c.alignment=Alignment(horizontal="center",wrap_text=True); c.border=Border(bottom=border); ws.column_dimensions[c.column_letter].width=widths[i-1]
    ws.row_dimensions[1].height=30
for s in ["Categories","Brands","Colors"]:
    if s in wb.sheetnames: del wb[s]
ws=wb.create_sheet("Categories",0); setup(ws,["category_name","category_fa_name","filter_name","icon"],[24,26,24,18])
for r in [["Mobile phones","گوشی موبایل","گوشی","📱"],["Tablets","تبلت","تبلت","▣"],["Laptops","لپ‌تاپ","لپ‌تاپ","💻"],["Accessories","لوازم جانبی","اکسسوری","🎧"]]: ws.append(r)
ws=wb.create_sheet("Brands",1); setup(ws,["brand_name","brand_fa_name","icon"],[22,24,18])
for r in [["Apple","اپل",""],["Samsung","سامسونگ","✦"],["Xiaomi","شیائومی","◉"],["Google","گوگل","G"]]: ws.append(r)
ws=wb.create_sheet("Colors",2); setup(ws,["color_name","color_fa_name"],[22,24])
for r in [["Black","مشکی"],["White","سفید"],["Blue","آبی"],["Orange","نارنجی"],["Green","سبز"]]: ws.append(r)
ws=wb["Products"]; old=list(ws.iter_rows(min_row=2,values_only=True))
heads=["product_id","Category","Brand","title","model","color","color_en","sku","price","old_price","discount%","stock","warranty","promotion","status","image_url","attribute_key","attribute_value"]
setup(ws,heads,[20,18,16,36,30,22,18,22,16,16,12,10,24,12,16,42,20,28])
for r in old: ws.append([r[0],r[14],r[15],r[1],r[2],r[3],r[4],r[5],r[6],r[7],r[8],r[9],r[10],bool(r[17]),r[12],r[13],"sim",r[16]])
for row in ws.iter_rows(min_row=2):
    for c in row: c.font=bf; c.alignment=Alignment(vertical="center",wrap_text=True); c.border=Border(bottom=border)
    row[8].number_format='#,##0'; row[9].number_format='#,##0'; row[10].number_format='0%'
for formula,rng in [("=Categories!$B$2:$B$100","B2:B1000"),("=Brands!$B$2:$B$100","C2:C1000"),("=Colors!$B$2:$B$100","F2:F1000"),('"بدون گارانتی,6 ماهه,12 ماهه,18 ماهه,24 ماهه"',"M2:M1000"),('"TRUE,FALSE"',"N2:N1000"),('"Active,Inactive,Draft"',"O2:O1000")]:
    dv=DataValidation(type="list",formula1=formula,allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
ws.conditional_formatting.add("L2:L1000",CellIsRule(operator="lessThan",formula=["1"],fill=PatternFill("solid",fgColor="FDE2E2")))
tab=Table(displayName="ProductsTable",ref=f"A1:R{max(ws.max_row,2)}"); tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True,showColumnStripes=False); ws.add_table(tab)
for s in ["Orders","Settings"]: wb[s].sheet_view.rightToLeft=True
wb.save("fast_order_catalog.xlsx")
